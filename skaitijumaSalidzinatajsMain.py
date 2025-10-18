#!/usr/bin/env python3
from __future__ import annotations

import datetime as dt
import json
import os
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import pandas as pd
from pandas import DataFrame
from rapidfuzz import fuzz, process
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None  # type: ignore[assignment]
    get_column_letter = None  # type: ignore[assignment]


def app_base_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS)  # type: ignore[attr-defined]
    return Path(__file__).resolve().parent


BASE_DIR = app_base_dir()
DEFAULT_SOURCE_DIR = BASE_DIR / "source"
DEFAULT_COUNTS_DIR = BASE_DIR / "counts"
DEFAULT_REPORTS_DIR = BASE_DIR / "reports"
CATALOG_XLSX = BASE_DIR / "allproducts.xlsx"
CATALOG_JS = BASE_DIR / "inventura" / "sources" / "products.js"
APP_CONFIG_NAME = "EziInventory"


def get_config_dir() -> Path:
    system = sys.platform
    if system.startswith("win"):
        base = os.environ.get("APPDATA")
        if base:
            return Path(base) / APP_CONFIG_NAME
        return Path.home() / "AppData" / "Roaming" / APP_CONFIG_NAME
    if system == "darwin":
        return Path.home() / "Library" / "Application Support" / APP_CONFIG_NAME
    base = os.environ.get("XDG_CONFIG_HOME", str(Path.home() / ".config"))
    return Path(base) / APP_CONFIG_NAME


CONFIG_FILE = get_config_dir() / "inventory_paths.json"

HEADER_ALIASES: Dict[str, Sequence[str]] = {
    "code": ("code", "kods", "sku", "artikuls", "preces kods", "produkts", "id"),
    "name": ("name", "nosaukums", "apraksts", "product", "title"),
    "amount": ("amount", "qty", "quantity", "count", "skaits", "daudzums", "inventura"),
}

COLUMN_ORDER = ("code", "name", "source_qty", "user_qty", "difference")


PDF_FONT_NAME = "EzuInventoryFont"
_pdf_font_registered = False
_pdf_font_active_name = "Helvetica"


def locate_font_file() -> Optional[Path]:
    preferred: List[Path] = []
    if sys.platform.startswith("win"):
        windir = Path(os.environ.get("WINDIR", "C:\\Windows")) / "Fonts"
        preferred.extend(
            [
                windir / "arialuni.ttf",
                windir / "segoeui.ttf",
                windir / "arial.ttf",
                windir / "calibri.ttf",
            ]
        )
    preferred.extend(
        [
            Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
            Path("/System/Library/Fonts/Supplemental/Times New Roman.ttf"),
            Path("/Library/Fonts/Arial Unicode.ttf"),
            Path("/Library/Fonts/Arial.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        ]
    )
    for candidate in preferred:
        if candidate.exists():
            return candidate
    search_dirs = [
        Path.home() / "Library" / "Fonts",
        Path("/Library/Fonts"),
        Path("/System/Library/Fonts"),
        Path("/System/Library/Fonts/Supplemental"),
    ]
    for directory in search_dirs:
        if not directory.exists():
            continue
        for candidate in sorted(directory.glob("*.ttf")):
            if candidate.exists():
                return candidate
    return None


def ensure_pdf_font_registered() -> str:
    global _pdf_font_registered, _pdf_font_active_name
    if _pdf_font_registered:
        return _pdf_font_active_name
    font_path = locate_font_file()
    if font_path is not None:
        try:
            pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, str(font_path)))
            _pdf_font_active_name = PDF_FONT_NAME
        except Exception:
            _pdf_font_active_name = "Helvetica"
    _pdf_font_registered = True
    return _pdf_font_active_name


def load_path_config() -> Dict[str, Path]:
    defaults = {
        "source": DEFAULT_SOURCE_DIR,
        "counts": DEFAULT_COUNTS_DIR,
        "reports": DEFAULT_REPORTS_DIR,
    }
    if CONFIG_FILE.exists():
        try:
            data = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
            for key in defaults:
                value = data.get(key)
                if isinstance(value, str) and value.strip():
                    path = Path(value).expanduser()
                    try:
                        defaults[key] = path.resolve()
                    except Exception:
                        defaults[key] = path
        except Exception:
            pass
    return defaults


def save_path_config(paths: Dict[str, Path]) -> None:
    payload = {key: str(path) for key, path in paths.items()}
    try:
        CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
        CONFIG_FILE.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    except Exception:
        pass


def normalize_header(name: str) -> Optional[str]:
    cleaned = re.sub(r"\s+", " ", name.strip().lower())
    for canonical, aliases in HEADER_ALIASES.items():
        if cleaned in aliases:
            return canonical
    return None


def normalize_number(value: object) -> Optional[float]:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("\u00a0", " ")
    text = text.replace(" ", "")
    text = text.replace(",", ".")
    match = re.search(r"[+-]?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def load_catalog() -> Dict[str, str]:
    catalog: Dict[str, str] = {}
    if CATALOG_XLSX.exists():
        df = pd.read_excel(CATALOG_XLSX, dtype=str)
        columns = {normalize_header(col): col for col in df.columns}
        code_col, name_col = columns.get("code"), columns.get("name")
        if code_col and name_col:
            for _, row in df[[code_col, name_col]].dropna().iterrows():
                code = str(row[code_col]).strip()
                name = str(row[name_col]).strip()
                if code:
                    catalog[code] = name
    if CATALOG_JS.exists():
        try:
            text = CATALOG_JS.read_text(encoding="utf-8")
            data = json.loads(re.search(r"=\s*(\{.*\})\s*;?\s*$", text, re.S).group(1))  # type: ignore[arg-type]
            for key, value in data.items():
                catalog.setdefault(str(key).strip(), str(value).strip())
        except Exception:
            pass
    return catalog


def load_table(path: Path) -> DataFrame:
    if path.suffix.lower() == ".csv":
        df = pd.read_csv(path, dtype=str)
    elif path.suffix.lower() == ".xlsx":
        df = pd.read_excel(path, dtype=str)
    else:
        raise ValueError(f"Neatbalstīts formāts: {path.suffix}")
    mapping = {}
    for column in df.columns:
        canonical = normalize_header(column)
        if canonical and canonical not in mapping:
            mapping[canonical] = column
    required = {"code", "amount"}
    if not required.issubset(mapping):
        raise ValueError(f"Kolonnas netika atpazītas: {path.name}")
    df = df.rename(columns={v: k for k, v in mapping.items()})
    df["code"] = df["code"].astype(str).str.strip()
    df = df[df["code"] != ""]
    if "name" not in df.columns:
        df["name"] = ""
    df["name"] = df["name"].fillna("").astype(str).str.strip()
    df["amount"] = df["amount"].apply(normalize_number)
    df = df.dropna(subset=["amount"])
    return df[["code", "name", "amount"]]


def aggregate_sources(paths: Sequence[Path], catalog: Dict[str, str]) -> DataFrame:
    frames = []
    for path in paths:
        df = load_table(path)
        if df.empty:
            continue
        df["name"] = df["name"].replace("", None)
        frames.append(df)
    if not frames:
        return pd.DataFrame(columns=["code", "name", "amount"])
    combined = pd.concat(frames, ignore_index=True)
    combined["name"] = combined["name"].fillna("")
    grouped = combined.groupby("code", as_index=False).agg(
        amount=("amount", "sum"),
        name=("name", lambda values: next((v for v in values if v), "")),
    )
    grouped["name"] = grouped.apply(
        lambda row: row["name"] or catalog.get(row["code"], ""), axis=1
    )
    return grouped[["code", "name", "amount"]]


def aggregate_counts(paths: Sequence[Path], catalog: Dict[str, str]) -> DataFrame:
    frames = []
    for path in paths:
        df = load_table(path)
        if df.empty:
            continue
        df["name"] = df["name"].replace("", None)
        frames.append(df)
    if not frames:
        return pd.DataFrame(columns=["code", "name", "count"])
    combined = pd.concat(frames, ignore_index=True)
    combined["name"] = combined["name"].fillna("")
    grouped = combined.groupby("code", as_index=False).agg(
        count=("amount", "sum"),
        name=("name", lambda values: next((v for v in values if v), "")),
    )
    grouped["name"] = grouped.apply(
        lambda row: row["name"] or catalog.get(row["code"], ""), axis=1
    )
    return grouped[["code", "name", "count"]]


def build_comparison(
    source_df: DataFrame,
    counts_df: DataFrame,
    catalog: Dict[str, str],
) -> DataFrame:
    prepared_source = source_df.rename(columns={"amount": "source_qty", "name": "name_source"})
    prepared_counts = counts_df.rename(columns={"count": "user_qty", "name": "name_count"})
    merged = prepared_source.merge(
        prepared_counts,
        on="code",
        how="outer",
        suffixes=("", "_count"),
    )
    merged["source_qty"] = merged["source_qty"].astype(float)
    merged["user_qty"] = merged["user_qty"].astype(float)
    if "name_source" in merged.columns:
        merged["name_source"] = merged["name_source"].fillna("")
    else:
        merged["name_source"] = ""
    if "name_count" in merged.columns:
        merged["name_count"] = merged["name_count"].fillna("")
    else:
        merged["name_count"] = ""
    merged["name"] = merged.apply(
        lambda row: row["name_source"]
        or row["name_count"]
        or catalog.get(row["code"], ""),
        axis=1,
    )
    merged["category"] = merged.apply(classify_row, axis=1)
    merged["difference"] = merged["user_qty"].fillna(0) - merged["source_qty"].fillna(0)
    merged["name"] = merged["name"].fillna("").astype(str)
    final = merged[["code", "name", "source_qty", "user_qty", "difference", "category"]]
    final = final.sort_values("code").reset_index(drop=True)
    return final


def classify_row(row: pd.Series) -> str:
    if pd.isna(row["source_qty"]):
        return "extra"
    if pd.isna(row["user_qty"]):
        return "missing"
    if abs(row["user_qty"] - row["source_qty"]) < 1e-9:
        return "match"
    return "partial"


def format_quantity(value: Optional[float], *, zero_if_none: bool = False) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "0" if zero_if_none else ""
    number = float(value)
    if abs(number - round(number)) < 1e-9:
        return f"{int(round(number))}"
    return f"{number:.2f}".rstrip("0").rstrip(".")


def format_difference(value: Optional[float]) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "0"
    number = float(value)
    if abs(number) < 1e-9:
        return "0"
    prefix = "+" if number > 0 else ""
    return f"{prefix}{format_quantity(number)}"


def export_to_pdf(path: Path, headers: Sequence[str], rows: Sequence[Sequence[str]], subtitle: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=30, rightMargin=30, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    font_name = ensure_pdf_font_registered()
    styles["Title"].fontName = font_name
    styles["Normal"].fontName = font_name
    story = []
    story.append(Paragraph("Inventāra salīdzinājums", styles["Title"]))
    story.append(Paragraph(subtitle, styles["Normal"]))
    story.append(Spacer(1, 12))

    table_data = [list(headers)]
    for row in rows:
        table_data.append([str(value) for value in row])

    table = Table(table_data, repeatRows=1)
    table_style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), font_name),
            ("FONTNAME", (0, 1), (-1, -1), font_name),
            ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.2, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ]
    )
    table.setStyle(table_style)
    story.append(table)
    doc.build(story)


def export_to_xlsx(
    path: Path,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    display_rows: Sequence[Sequence[str]],
) -> None:
    if Workbook is None or get_column_letter is None:
        raise RuntimeError("Nav instalēts 'openpyxl'. Uzstādi ar 'pip install openpyxl'.")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Pārskats"
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append(list(row))
    widths = [len(str(header)) for header in headers]
    for row in display_rows:
        for idx, value in enumerate(row):
            widths[idx] = max(widths[idx], len(str(value)))
    for idx, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(idx)].width = min(60, width + 2)
    workbook.save(path)


@dataclass
class CatalogItem:
    code: str
    name: str

    @property
    def display(self) -> str:
        return f"{self.code} – {self.name}"


class SearchDialog(tk.Toplevel):
    def __init__(self, parent: "InventoryApp"):
        super().__init__(parent)
        self.parent = parent
        self.title("Produktu meklētājs")
        self.resizable(False, False)
        self.protocol("WM_DELETE_WINDOW", self.close)
        self.query_var = tk.StringVar()
        self.result_items: List[CatalogItem] = []

        frame = ttk.Frame(self, padding=10)
        frame.pack(fill="both", expand=True)

        ttk.Label(frame, text="Ievadi kodu vai nosaukumu:").pack(anchor="w")
        entry = ttk.Entry(frame, textvariable=self.query_var, width=40)
        entry.pack(fill="x", pady=(4, 8))
        entry.focus_set()

        list_frame = ttk.Frame(frame)
        list_frame.pack(fill="both", expand=True)
        self.listbox = tk.Listbox(list_frame, height=15)
        self.listbox.pack(side="left", fill="both", expand=True)
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.listbox.yview)
        scrollbar.pack(side="right", fill="y")
        self.listbox.configure(yscrollcommand=scrollbar.set)

        entry.bind("<KeyRelease>", lambda _e: self.refresh())
        self.listbox.bind("<Double-1>", lambda _e: self.activate_selection())
        self.listbox.bind("<Return>", lambda _e: self.activate_selection())

        self.refresh()
        self.grab_set()

    def refresh(self) -> None:
        query = self.query_var.get()
        self.result_items = self.parent.search_catalog(query, limit=30)
        self.listbox.delete(0, tk.END)
        for item in self.result_items:
            self.listbox.insert(tk.END, item.display)
        if self.result_items:
            self.listbox.selection_set(0)

    def activate_selection(self) -> None:
        selection = self.listbox.curselection()
        if not selection:
            return
        item = self.result_items[selection[0]]
        self.parent.focus_code(item.code)

    def close(self) -> None:
        self.parent.search_dialog = None
        self.destroy()


class SettingsDialog(tk.Toplevel):
    def __init__(self, parent: "InventoryApp"):
        super().__init__(parent)
        self.parent = parent
        self.title("Mapju iestatījumi")
        self.resizable(False, False)
        self.protocol("WM_DELETE_WINDOW", self.close)
        self.transient(parent)

        self.source_var = tk.StringVar(value=str(parent.source_dir))
        self.counts_var = tk.StringVar(value=str(parent.counts_dir))
        self.reports_var = tk.StringVar(value=str(parent.reports_dir))

        frame = ttk.Frame(self, padding=10)
        frame.pack(fill="both", expand=True)

        self._build_row(frame, 0, "Avota mape:", self.source_var)
        self._build_row(frame, 1, "Skaitījumu mape:", self.counts_var)
        self._build_row(frame, 2, "Eksporta mape:", self.reports_var)

        button_row = ttk.Frame(frame)
        button_row.grid(row=3, column=0, columnspan=3, pady=(10, 0), sticky="e")
        ttk.Button(button_row, text="Atcelt", command=self.close).pack(side="right", padx=5)
        ttk.Button(button_row, text="Saglabāt", command=self.save).pack(side="right")

        self.grab_set()
        self.wait_visibility()
        self.focus()

    def _build_row(self, parent: ttk.Frame, row: int, label: str, variable: tk.StringVar) -> None:
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", pady=4)
        entry = ttk.Entry(parent, textvariable=variable, width=60)
        entry.grid(row=row, column=1, sticky="ew", padx=5)
        ttk.Button(parent, text="Izvēlēties…", command=lambda: self._browse(variable)).grid(row=row, column=2, sticky="w")
        parent.columnconfigure(1, weight=1)

    def _browse(self, variable: tk.StringVar) -> None:
        initial = variable.get() or str(BASE_DIR)
        selected = filedialog.askdirectory(initialdir=initial)
        if selected:
            variable.set(selected)

    def save(self) -> None:
        source = Path(self.source_var.get()).expanduser()
        counts = Path(self.counts_var.get()).expanduser()
        reports = Path(self.reports_var.get()).expanduser()
        if self.parent.apply_directory_settings(source, counts, reports):
            self.close()

    def close(self) -> None:
        self.parent.settings_dialog = None
        self.destroy()


class InventoryApp(tk.Tk):
    TAB_LABELS = {
        "match": "Precīzi sakrīt",
        "partial": "Daļēji sakrīt",
        "missing": "Nav atrasts",
        "extra": "Papildu",
    }

    def __init__(self) -> None:
        super().__init__()
        self.title("EŽI inventūras salīdzinātājs")
        self.geometry("1080x700")
        self.catalog = load_catalog()
        path_config = load_path_config()
        self.source_dir = path_config["source"]
        self.counts_dir = path_config["counts"]
        self.reports_dir = path_config["reports"]
        self.source_dir.mkdir(parents=True, exist_ok=True)
        self.counts_dir.mkdir(parents=True, exist_ok=True)
        self.reports_dir.mkdir(parents=True, exist_ok=True)
        self.path_config = path_config
        self.catalog_items = [CatalogItem(code=k, name=v) for k, v in sorted(self.catalog.items())]
        self.source_files = self._list_inventory_files(self.source_dir)
        self.count_files = self._list_inventory_files(self.counts_dir)
        self.comparison_df: DataFrame = pd.DataFrame(columns=["code", "name", "source_qty", "user_qty", "difference", "category"])
        self.search_dialog: Optional[SearchDialog] = None
        self.settings_dialog: Optional["SettingsDialog"] = None

        self._build_layout()
        self._refresh_file_lists()

    def _build_layout(self) -> None:
        main = ttk.Frame(self)
        main.pack(fill="both", expand=True)

        selection = ttk.Frame(main, padding=10)
        selection.pack(fill="x")

        source_frame = ttk.LabelFrame(selection, text="Avota faili (source/)")
        source_frame.pack(side="left", fill="both", expand=True, padx=(0, 10))
        source_inner = ttk.Frame(source_frame)
        source_inner.pack(fill="both", expand=True, padx=6, pady=6)
        self.source_listbox = tk.Listbox(source_inner, height=8, selectmode=tk.MULTIPLE, exportselection=False)
        self.source_listbox.pack(side="left", fill="both", expand=True)
        source_scroll = ttk.Scrollbar(source_inner, orient="vertical", command=self.source_listbox.yview)
        source_scroll.pack(side="right", fill="y")
        self.source_listbox.configure(yscrollcommand=source_scroll.set)

        count_frame = ttk.LabelFrame(selection, text="Skaitījuma faili (counts/)")
        count_frame.pack(side="left", fill="both", expand=True)
        count_inner = ttk.Frame(count_frame)
        count_inner.pack(fill="both", expand=True, padx=6, pady=6)
        self.count_listbox = tk.Listbox(count_inner, height=8, selectmode=tk.MULTIPLE)
        self.count_listbox.pack(side="left", fill="both", expand=True)
        count_scroll = ttk.Scrollbar(count_inner, orient="vertical", command=self.count_listbox.yview)
        count_scroll.pack(side="right", fill="y")
        self.count_listbox.configure(yscrollcommand=count_scroll.set)

        button_bar = ttk.Frame(main, padding=(10, 0))
        button_bar.pack(fill="x")
        ttk.Button(button_bar, text="Pārlādēt failus", command=self._refresh_file_lists).pack(side="left")
        ttk.Button(button_bar, text="Salīdzināt", command=self.perform_comparison).pack(side="left", padx=5)
        ttk.Button(button_bar, text="Meklēt produktu", command=self.open_search_dialog).pack(side="left", padx=5)
        self.match_button = ttk.Button(button_bar, text="Atzīmēt kā sakrītu", command=self.mark_as_matched, state="disabled")
        self.match_button.pack(side="left", padx=5)
        ttk.Button(button_bar, text="Eksportēt", command=self.export_results).pack(side="left", padx=5)
        ttk.Button(button_bar, text="Iestatījumi", command=self.open_settings_dialog).pack(side="left", padx=5)

        self.notebook = ttk.Notebook(main)
        self.notebook.pack(fill="both", expand=True, padx=10, pady=10)
        self.notebook.bind("<<NotebookTabChanged>>", lambda _e: self._update_match_button())

        self.trees: Dict[str, ttk.Treeview] = {}
        for key, title in self.TAB_LABELS.items():
            frame = ttk.Frame(self.notebook)
            self.notebook.add(frame, text=title)
            tree = ttk.Treeview(
                frame,
                columns=("code", "name", "source", "user", "diff"),
                show="headings",
                selectmode="extended",
            )
            tree.heading("code", text="Kods")
            tree.heading("name", text="Nosaukums")
            tree.heading("source", text="Avota daudzums")
            tree.heading("user", text="Skaitītais daudzums")
            tree.heading("diff", text="Atšķirība")
            tree.column("code", width=140, anchor="w")
            tree.column("name", width=420, anchor="w")
            tree.column("source", width=120, anchor="e")
            tree.column("user", width=120, anchor="e")
            tree.column("diff", width=120, anchor="e")
            tree.pack(side="left", fill="both", expand=True)
            scrollbar = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
            scrollbar.pack(side="right", fill="y")
            tree.configure(yscrollcommand=scrollbar.set)
            tree.bind("<Double-1>", lambda event, category=key: self.start_edit(event, category))
            tree.bind("<<TreeviewSelect>>", lambda _e: self._update_match_button())
            self.trees[key] = tree

    def _list_inventory_files(self, folder: Path) -> List[Path]:
        if not folder.exists():
            return []
        return sorted(f for f in folder.iterdir() if f.suffix.lower() in {".csv", ".xlsx"} and f.is_file())

    def _refresh_file_lists(self) -> None:
        self.source_files = self._list_inventory_files(self.source_dir)
        self.count_files = self._list_inventory_files(self.counts_dir)
        self.source_listbox.delete(0, tk.END)
        for path in self.source_files:
            self.source_listbox.insert(tk.END, path.name)
        self.count_listbox.delete(0, tk.END)
        for path in self.count_files:
            self.count_listbox.insert(tk.END, path.name)

    def open_settings_dialog(self) -> None:
        if self.settings_dialog and self.settings_dialog.winfo_exists():
            self.settings_dialog.lift()
            self.settings_dialog.focus_force()
            return
        self.settings_dialog = SettingsDialog(self)

    def apply_directory_settings(self, source: Path, counts: Path, reports: Path) -> bool:
        errors: List[str] = []

        def prepare(path: Path, label: str) -> Optional[Path]:
            try:
                resolved = path.expanduser()
                if resolved.exists() and not resolved.is_dir():
                    errors.append(f"{label}: norādītais ceļš nav mape.")
                    return None
                resolved.mkdir(parents=True, exist_ok=True)
                return resolved.resolve()
            except Exception as exc:
                errors.append(f"{label}: {exc}")
                return None

        prepared_source = prepare(source, "Avota mape")
        prepared_counts = prepare(counts, "Skaitījumu mape")
        prepared_reports = prepare(reports, "Eksporta mape")

        if errors or not all([prepared_source, prepared_counts, prepared_reports]):
            messagebox.showerror("Uzstādījumi", "\n".join(errors))
            return False

        self.source_dir = prepared_source  # type: ignore[assignment]
        self.counts_dir = prepared_counts  # type: ignore[assignment]
        self.reports_dir = prepared_reports  # type: ignore[assignment]
        self.reports_dir.mkdir(parents=True, exist_ok=True)

        self.path_config = {
            "source": self.source_dir,
            "counts": self.counts_dir,
            "reports": self.reports_dir,
        }
        save_path_config(self.path_config)
        self._refresh_file_lists()
        messagebox.showinfo("Uzstādījumi", "Mapju uzstādījumi saglabāti.")
        return True

    def open_search_dialog(self) -> None:
        if self.search_dialog and self.search_dialog.winfo_exists():
            self.search_dialog.lift()
            self.search_dialog.focus_force()
            return
        self.search_dialog = SearchDialog(self)

    def search_catalog(self, query: str, limit: int = 20) -> List[CatalogItem]:
        query = query.strip()
        if not query:
            return self.catalog_items[:limit]
        if query.isdigit():
            filtered = [item for item in self.catalog_items if item.code.startswith(query)]
            return filtered[:limit]
        choices = [item.display for item in self.catalog_items]
        results = process.extract(query, choices, scorer=fuzz.WRatio, limit=limit)
        return [self.catalog_items[idx] for _match, score, idx in results if score > 30]

    def perform_comparison(self) -> None:
        source_selection = self.source_listbox.curselection()
        if not source_selection:
            messagebox.showerror("Salīdzinājums", "Izvēlies vismaz vienu avota failu.")
            return
        source_paths = [self.source_files[i] for i in source_selection]
        count_indices = self.count_listbox.curselection()
        if not count_indices:
            messagebox.showerror("Salīdzinājums", "Izvēlies vismaz vienu skaitījuma failu.")
            return
        count_paths = [self.count_files[i] for i in count_indices]
        try:
            source_df = aggregate_sources(source_paths, self.catalog)
            if source_df.empty:
                raise ValueError("Avota failos nav derīgu ierakstu.")
            counts_df = aggregate_counts(count_paths, self.catalog)
            self.comparison_df = build_comparison(source_df, counts_df, self.catalog)
        except Exception as exc:
            messagebox.showerror("Salīdzinājums", str(exc))
            return
        self.populate_trees()
        messagebox.showinfo("Salīdzinājums", "Salīdzinājums pabeigts.")

    def populate_trees(self) -> None:
        for category, tree in self.trees.items():
            tree.delete(*tree.get_children())
            subset = self.comparison_df[self.comparison_df["category"] == category]
            for _, row in subset.iterrows():
                values = (
                    row["code"],
                    row["name"],
                    format_quantity(row["source_qty"]),
                    format_quantity(row["user_qty"]),
                    format_difference(row["difference"]),
                )
                tags = ()
                if category == "partial":
                    diff = row["difference"]
                    if diff is not None and not pd.isna(diff):
                        if diff > 0:
                            tags = ("above",)
                        elif diff < 0:
                            tags = ("below",)
                tree.insert("", "end", iid=row["code"], values=values, tags=tags)
            tree.tag_configure("above", background="#c8e6c9")
            tree.tag_configure("below", background="#ffcdd2")
        self._update_match_button()

    def start_edit(self, event: tk.Event, category: str) -> None:
        tree = self.trees[category]
        item_id = tree.identify_row(event.y)
        column = tree.identify_column(event.x)
        if not item_id or column != "#4":
            return
        bbox = tree.bbox(item_id, column)
        if not bbox:
            return
        x, y, width, height = bbox
        current_value = tree.set(item_id, "user")
        entry = tk.Entry(tree)
        entry.insert(0, current_value)
        entry.place(x=x, y=y, width=width, height=height)
        entry.focus_set()

        def commit(_event: Optional[tk.Event] = None) -> None:
            text = entry.get().strip()
            entry.destroy()
            if text == "":
                new_value = None
            else:
                new_value = normalize_number(text)
                if new_value is None:
                    messagebox.showerror("Rediģēšana", "Nederīgs daudzums.")
                    return
            self.update_user_quantity(item_id, new_value)

        entry.bind("<Return>", commit)
        entry.bind("<FocusOut>", commit)

    def update_user_quantity(self, code: str, value: Optional[float]) -> None:
        if self.comparison_df.empty:
            return
        mask = self.comparison_df["code"] == code
        if not mask.any():
            return
        self.comparison_df.loc[mask, "user_qty"] = value
        user_series = self.comparison_df.loc[mask, "user_qty"].fillna(0)
        source_series = self.comparison_df.loc[mask, "source_qty"].fillna(0)
        self.comparison_df.loc[mask, "difference"] = user_series - source_series
        self.comparison_df.loc[mask, "category"] = self.comparison_df.loc[mask].apply(classify_row, axis=1)
        self.populate_trees()

    def mark_as_matched(self) -> None:
        tab_key = self._current_tab()
        tree = self.trees[tab_key]
        for code in tree.selection():
            mask = self.comparison_df["code"] == code
            if not mask.any():
                continue
            source_qty = self.comparison_df.loc[mask, "source_qty"]
            self.comparison_df.loc[mask, "user_qty"] = source_qty
            self.comparison_df.loc[mask, "difference"] = 0.0
            self.comparison_df.loc[mask, "category"] = "match"
        self.populate_trees()

    def focus_code(self, code: str) -> None:
        if self.comparison_df.empty:
            messagebox.showinfo("Meklēšana", "Salīdzinājums nav ielādēts.")
            return
        row = self.comparison_df[self.comparison_df["code"] == code]
        if row.empty:
            messagebox.showinfo("Meklēšana", "Produkts nav pašreizējā salīdzinājumā.")
            return
        category = row.iloc[0]["category"]
        tab_index = list(self.TAB_LABELS.keys()).index(category)
        self.notebook.select(tab_index)
        tree = self.trees[category]
        tree.selection_set(code)
        tree.see(code)
        self._update_match_button()

    def export_results(self) -> None:
        if self.comparison_df.empty:
            messagebox.showerror("Eksports", "Nav datu ko eksportēt.")
            return
        dialog = ExportDialog(self)
        self.wait_window(dialog)
        if not dialog.result:
            return
        scope = dialog.result["scope"]
        export_format = dialog.result["format"]
        columns = dialog.result["columns"]
        label_map = {
            "code": "Kods",
            "name": "Nosaukums",
            "source_qty": "Avota daudzums",
            "user_qty": "Skaitītais daudzums",
            "difference": "Atšķirība",
        }
        if scope == "diff":
            export_df = self.comparison_df[self.comparison_df["category"] != "match"].copy()
        else:
            export_df = self.comparison_df.copy()
        column_keys = list(columns)
        headers = [label_map[col] for col in column_keys]

        numeric_rows: List[List[object]] = []
        display_rows: List[List[str]] = []
        for _, row in export_df[column_keys].iterrows():
            numeric_row: List[object] = []
            display_row: List[str] = []
            for col in column_keys:
                value = row[col]
                if col in {"source_qty", "user_qty", "difference"}:
                    numeric_value: float = 0.0
                    if not pd.isna(value):
                        try:
                            numeric_value = float(value)  # type: ignore[arg-type]
                        except (TypeError, ValueError):
                            normalized = normalize_number(value)
                            if normalized is not None:
                                numeric_value = float(normalized)
                    numeric_row.append(numeric_value)
                    if col == "difference":
                        display_row.append(format_difference(numeric_value))
                    else:
                        display_row.append(format_quantity(numeric_value, zero_if_none=True))
                else:
                    text_value = "" if pd.isna(value) else str(value)
                    numeric_row.append(text_value)
                    display_row.append(text_value)
            numeric_rows.append(numeric_row)
            display_rows.append(display_row)

        timestamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
        subtitle = f"Izveidots: {timestamp}"
        try:
            if export_format == "xlsx":
                path = self.reports_dir / f"salidzinajums-{timestamp}.xlsx"
                export_to_xlsx(path, headers, numeric_rows, display_rows)
            else:
                path = self.reports_dir / f"salidzinajums-{timestamp}.pdf"
                export_to_pdf(path, headers, display_rows, subtitle)
        except Exception as exc:
            messagebox.showerror("Eksports", f"Neizdevās izveidot failu: {exc}")
            return
        flash_fn = getattr(self, "flash", None)
        if callable(flash_fn):
            try:
                flash_fn()
            except tk.TclError:
                pass
        try:
            self.bell()
        except (tk.TclError, AttributeError):
            pass
        messagebox.showinfo("Eksports", f"Fails saglabāts: {path.name}")

    def _current_tab(self) -> str:
        tab_index = self.notebook.index(self.notebook.select())
        return list(self.TAB_LABELS.keys())[tab_index]

    def _update_match_button(self) -> None:
        tab = self._current_tab()
        tree = self.trees[tab]
        enabled = False
        for item in tree.selection():
            mask = self.comparison_df["code"] == item
            if mask.any():
                value = self.comparison_df.loc[mask, "source_qty"].iloc[0]
                if not pd.isna(value):
                    enabled = True
                    break
        self.match_button.configure(state="normal" if enabled else "disabled")


class ExportDialog(tk.Toplevel):
    def __init__(self, parent: InventoryApp):
        super().__init__(parent)
        self.title("Eksporta opcijas")
        self.resizable(False, False)
        self.result: Optional[Dict[str, object]] = None

        scope_frame = ttk.LabelFrame(self, text="Dati")
        scope_frame.grid(row=0, column=0, padx=10, pady=5, sticky="ew")
        self.scope_var = tk.StringVar(value="diff")
        ttk.Radiobutton(scope_frame, text="Tikai atšķirības", variable=self.scope_var, value="diff").grid(row=0, column=0, sticky="w")
        ttk.Radiobutton(scope_frame, text="Visas rindas", variable=self.scope_var, value="all").grid(row=1, column=0, sticky="w")

        format_frame = ttk.LabelFrame(self, text="Formāts")
        format_frame.grid(row=1, column=0, padx=10, pady=5, sticky="ew")
        self.format_var = tk.StringVar(value="xlsx")
        ttk.Radiobutton(format_frame, text="XLSX", variable=self.format_var, value="xlsx").grid(row=0, column=0, sticky="w")
        ttk.Radiobutton(format_frame, text="PDF", variable=self.format_var, value="pdf").grid(row=1, column=0, sticky="w")

        columns_frame = ttk.LabelFrame(self, text="Kolonnas")
        columns_frame.grid(row=2, column=0, padx=10, pady=5, sticky="ew")
        self.column_vars: Dict[str, tk.BooleanVar] = {}
        for idx, (key, label) in enumerate(
            (
                ("code", "Kods"),
                ("name", "Nosaukums"),
                ("source_qty", "Avota daudzums"),
                ("user_qty", "Skaitītais daudzums"),
                ("difference", "Atšķirība"),
            )
        ):
            var = tk.BooleanVar(value=True if key in {"code", "name", "source_qty", "user_qty"} else False)
            self.column_vars[key] = var
            ttk.Checkbutton(columns_frame, text=label, variable=var).grid(row=idx, column=0, sticky="w")

        button_bar = ttk.Frame(self)
        button_bar.grid(row=3, column=0, padx=10, pady=10, sticky="e")
        ttk.Button(button_bar, text="Atcelt", command=self.destroy).grid(row=0, column=0, padx=5)
        ttk.Button(button_bar, text="Eksportēt", command=self._submit).grid(row=0, column=1)
        self.grab_set()
        self.wait_visibility()
        self.focus()

    def _submit(self) -> None:
        columns = [name for name, var in self.column_vars.items() if var.get()]
        if not columns:
            messagebox.showerror("Eksports", "Izvēlies vismaz vienu kolonnu.")
            return
        self.result = {"scope": self.scope_var.get(), "format": self.format_var.get(), "columns": columns}
        self.destroy()


def main() -> None:
    app = InventoryApp()
    app.mainloop()


if __name__ == "__main__":
    main()
